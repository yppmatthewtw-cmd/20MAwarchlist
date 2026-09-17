# -*- coding: utf-8 -*-
"""Independent verification of the 加息及地緣政治 3 日觀察清單 workbook.

Recomputes every cached value straight from the source JSON (never from the
builder's own code path) and checks every ticker cell carries the right link.
"""
import json, re
from openpyxl import load_workbook

PX  = json.load(open('data/ratehike_geo_px.json'))
SEC = json.load(open('data/ratehike_geo_sec.json'))
IDX = json.load(open('data/ratehike_geo_idx.json'))
FN  = 'reports/RateHike_Geopolitical_3Day_Watchlist_R1.01_claudeopus5high_09.17_1516.xlsx'
TV  = "https://www.tradingview.com/chart/Q1c5VWwD/?symbol={}"

V  = load_workbook(FN, data_only=True)   # cached values
Fm = load_workbook(FN)                   # formulas + hyperlinks
P  = []
def bad(m): P.append(m); print("  PROBLEM:", m)
def close(a, b, tol=1e-9):
    return a is not None and b is not None and abs(a - b) <= tol * max(1, abs(b))

print("sheets:", V.sheetnames)
LINKED = {}          # (sheet, coord) -> ticker, every cell we expect to be a link

def link_ok(sheet, coord, tk):
    c = Fm[sheet][coord]
    if c.value != tk:
        bad(f"{sheet}!{coord} value {c.value!r} != {tk!r}"); return
    if c.hyperlink is None:
        bad(f"{sheet}!{coord} ticker {tk} has no hyperlink"); return
    if c.hyperlink.target != TV.format(tk.lower()):
        bad(f"{sheet}!{coord} {tk} -> {c.hyperlink.target}"); return
    if not (c.font.underline and c.font.color and c.font.color.rgb.endswith("0563C1")):
        bad(f"{sheet}!{coord} {tk} not styled as a link")
    LINKED[(sheet, coord)] = tk

# 0. no error strings, every formula carries a cached value
errs = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
nf = nnone = 0
for wsF, wsV in zip(Fm.worksheets, V.worksheets):
    for rowF, rowV in zip(wsF.iter_rows(), wsV.iter_rows()):
        for cF, cV in zip(rowF, rowV):
            if isinstance(cV.value, str) and cV.value in errs:
                bad(f"{wsV.title}!{cV.coordinate} = {cV.value}")
            if isinstance(cF.value, str) and cF.value.startswith('='):
                nf += 1
                if cV.value is None:
                    nnone += 1; bad(f"{wsV.title}!{cF.coordinate} formula has no cached value")
print(f"formula cells: {nf}, without cached value: {nnone}")

# 1. Watchlist
WS = '3日受惠清單 Watchlist'
w = V[WS]
wts = [w.cell(6, c).value for c in (13, 14, 15, 16)]
if abs(sum(wts) - 1.0) > 1e-9: bad(f"weights sum {sum(wts)}")
print("weights:", wts, "sum cell Q6 =", w.cell(6, 17).value)
n, r = 0, 9
while w.cell(r, 3).value:
    tk = w.cell(r, 3).value; p = PX[tk]; n += 1
    link_ok(WS, f"C{r}", tk)
    if not close(w.cell(r, 5).value, p['c16']):                bad(f"{tk} close")
    if not close(w.cell(r, 8).value,  p['c16']/p['c15'] - 1):  bad(f"{tk} d16%")
    if not close(w.cell(r, 9).value,  p['c16']/p['c09'] - 1):  bad(f"{tk} 5d%")
    if not close(w.cell(r, 12).value, p['v16']/p['vmed']):     bad(f"{tk} rvol")
    m, fl, fr, bd = (w.cell(r, c).value for c in (13, 14, 15, 16))
    exp = (m*wts[0] + fl*wts[1] + fr*wts[2] + bd*wts[3]) * 20
    if not close(w.cell(r, 17).value, exp): bad(f"{tk} composite {w.cell(r,17).value} != {exp}")
    if not 0 <= w.cell(r, 17).value <= 100: bad(f"{tk} composite out of range")
    r += 1
print(f"watchlist rows verified: {n}")
avg = sum(w.cell(x, 17).value for x in range(9, 9+n)) / n
if not close(w.cell(9+n, 17).value, avg): bad("watchlist footer avg")
WTK = [w.cell(x, 3).value for x in range(9, 9+n)]

# 2. Avoid
AS = '迴避清單 Avoid'
a = V[AS]; r, m = 9, 0
while a.cell(r, 2).value:
    tk = a.cell(r, 2).value; p = PX[tk]; m += 1
    link_ok(AS, f"B{r}", tk)
    if not close(a.cell(r, 7).value,  p['c16']/p['c15'] - 1): bad(f"avoid {tk} d16%")
    if not close(a.cell(r, 8).value,  p['c16']/p['c09'] - 1): bad(f"avoid {tk} 5d%")
    if not close(a.cell(r, 11).value, p['v16']/p['vmed']):    bad(f"avoid {tk} rvol")
    r += 1
print(f"avoid rows verified: {m}")
ATK = [a.cell(x, 2).value for x in range(9, 9+m)]

# 3. Evidence: 111 rows, basket split one ticker per cell, all linked
ES = '板塊資金流證據 Evidence'
e = V[ES]
MAXB, NOTE_COL = 8, 23
nbask = 0
for i, s in enumerate(SEC['rows'], 1):
    rr = 8 + i
    if e.cell(rr, 1).value != i:       bad(f"evidence rank row {rr}")
    if e.cell(rr, 2).value != s['zh']: bad(f"evidence name row {rr}")
    if not close(e.cell(rr, 5).value, s['s5']): bad(f"evidence s5 {s['zh']}")
    if s['zh'] != 'LNG液化天然氣' and not close(e.cell(rr, 8).value, s['r5']):
        bad(f"evidence r5 {s['zh']}")
    bask = s['tick'].split(',')
    if len(bask) > MAXB: bad(f"{s['zh']} basket {len(bask)} > {MAXB}")
    for j, tk in enumerate(bask):
        link_ok(ES, f"{chr(ord('A')+14+j)}{rr}", tk); nbask += 1
    for j in range(len(bask), MAXB):
        if e.cell(rr, 15+j).value is not None:
            bad(f"{s['zh']} stray value in basket slot {j+1}")
    if e.cell(rr, NOTE_COL).value is None and s['zh'] in ('煉油與成品油', '頁岩E&P勘探生產'):
        bad(f"{s['zh']} note missing at col {NOTE_COL}")
if e.cell(8 + len(SEC['rows']) + 1, 1).value is not None: bad("evidence has extra row")
lng = 9 + next(i for i, x in enumerate(SEC['rows']) if x['zh'] == 'LNG液化天然氣')
if e.cell(lng, 8).value != 'n/a（數據缺陷）': bad("LNG row not flagged n/a")
print(f"evidence rows verified: {len(SEC['rows'])}, basket ticker links: {nbask}")

# 4. Ticker Index: every distinct ticker in the workbook, linked, tagged, priced
IS = '代號索引 Ticker Index'
ix = V[IS]
SUB = {}
for i, s in enumerate(SEC['rows'], 1):
    for t in s['tick'].split(','):
        SUB.setdefault(t, []).append((i, s['zh']))
seen, r = [], 9
while ix.cell(r, 1).value:
    tk = ix.cell(r, 1).value; seen.append(tk)
    link_ok(IS, f"A{r}", tk)
    p = IDX[tk]
    if not close(ix.cell(r, 6).value,  p['c16']):               bad(f"index {tk} close")
    if not close(ix.cell(r, 9).value,  p['c16']/p['c15'] - 1):  bad(f"index {tk} d16%")
    if not close(ix.cell(r, 10).value, p['c16']/p['c09'] - 1):  bad(f"index {tk} 5d%")
    if not close(ix.cell(r, 13).value, p['v16']/p['vmed']):     bad(f"index {tk} rvol")
    exp_w = next((t for t, x in zip([w.cell(y,1).value for y in range(9,9+n)], WTK) if x == tk), "—")
    if ix.cell(r, 2).value != exp_w: bad(f"index {tk} watchlist tag {ix.cell(r,2).value} != {exp_w}")
    if (ix.cell(r, 3).value != "—") != (tk in ATK): bad(f"index {tk} avoid tag wrong")
    subs = sorted(SUB.get(tk, []))
    exp_s = "、".join(z for _, z in subs) or "—"
    if ix.cell(r, 4).value != exp_s: bad(f"index {tk} sub-sectors {ix.cell(r,4).value!r} != {exp_s!r}")
    if ix.cell(r, 5).value != (subs[0][0] if subs else None): bad(f"index {tk} best rank")
    r += 1
if seen != sorted(IDX): bad(f"index not the full sorted universe ({len(seen)} rows vs {len(IDX)})")
print(f"index rows verified: {len(seen)}")

# 4b. Catalysts sheet: every 相關代號 cell linked
CS = '催化劑日程 Catalysts'
cs = V[CS]; ncat = 0
r = 9
while cs.cell(r, 1).value:
    gap = False
    for j in range(7):
        tk = cs.cell(r, 8 + j).value
        if tk is None:
            gap = True; continue
        if gap: bad(f"{CS} row {r}: ticker at slot {j+1} after an empty slot")
        link_ok(CS, f"{chr(ord('A')+7+j)}{r}", tk); ncat += 1
    r += 1
print(f"catalyst ticker links: {ncat}")

# 5. every ticker cell in the workbook is linked, and every one is in the index
expect = n + m + nbask + len(seen) + ncat
if len(LINKED) != expect: bad(f"linked cells {len(LINKED)} != expected {expect}")
univ = set(WTK) | set(ATK) | {t for s in SEC['rows'] for t in s['tick'].split(',')}
if univ - set(seen): bad(f"tickers missing from index: {sorted(univ - set(seen))[:10]}")
nolink = []
for ws in Fm.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if (isinstance(c.value, str) and re.fullmatch(r"[A-Z]{1,5}(\.[A-Z])?", c.value)
                    and c.value in univ and (ws.title, c.coordinate) not in LINKED):
                nolink.append(f"{ws.title}!{c.coordinate}={c.value}")
if nolink: bad(f"ticker cells without a link: {nolink[:10]}")
print(f"ticker hyperlinks: {len(LINKED)} (watchlist {n} + avoid {m} + baskets {nbask} + index {len(seen)} + catalysts {ncat})")

# 6. Scenarios + Summary
s = V['情境矩陣 Scenarios']
pa = s['E4'].value
if not close(s['E5'].value, 1 - pa): bad("P(B)")
hdr = next(rr for rr in range(1, 30) if s.cell(rr, 1).value == '組別')
r, cnt = hdr + 1, 0
while s.cell(r, 2).value and s.cell(r, 4).value is not None:
    d, ee, f = s.cell(r, 4).value, s.cell(r, 5).value, s.cell(r, 6).value
    if not close(f, pa*d + (1-pa)*ee): bad(f"scenario row {r}")
    cnt += 1; r += 1
print(f"scenario rows verified: {cnt} (P(A)={pa}, header row {hdr})")
if cnt != 10: bad(f"scenario row count {cnt} != 10")
if not close(s.cell(r,6).value, sum(s.cell(x,6).value for x in range(hdr+1,hdr+7))/6): bad("scenario benefit avg")
if not close(s.cell(r+1,6).value, sum(s.cell(x,6).value for x in range(hdr+7,hdr+11))/4): bad("scenario avoid avg")

su = V['摘要 Summary']; found = {}
for rr in range(1, 40):
    for cc in (1, 4):
        lab = su.cell(rr, cc).value
        if isinstance(lab, str) and lab in ('受惠股檔數','迴避股檔數','受惠股平均綜合分','綜合分 ≥ 80 檔數'):
            found[lab] = su.cell(rr, cc+1).value
if found.get('受惠股檔數') != n: bad(f"summary count {found.get('受惠股檔數')} != {n}")
if found.get('迴避股檔數') != m: bad(f"summary avoid count {found.get('迴避股檔數')} != {m}")
if not close(found.get('受惠股平均綜合分'), avg): bad("summary avg composite")
exp80 = sum(1 for x in range(9, 9+n) if w.cell(x, 17).value >= 80)
if found.get('綜合分 ≥ 80 檔數') != exp80: bad(f"summary >=80 != {exp80}")
print("summary stats:", found)

import pandas as pd
for sh in V.sheetnames:
    if pd.read_excel(FN, sheet_name=sh, header=None).empty: bad(f"pandas read {sh} empty")
print("pandas read all sheets OK")
print("\nPROBLEMS:", len(P))
