# -*- coding: utf-8 -*-
"""Independent check of a money-flow watchlist workbook against its flow JSON.

Every number in these workbooks is an engine output rendered into a cell, so the check is
that the rendering is faithful and complete: no row lost, no column shifted, no ticker
without a link, and every figure equal to the source.

    python3 scripts/verify_flow_workbook1.py sub | ai
"""
import json, sys
from openpyxl import load_workbook

MODE = (sys.argv[1] if len(sys.argv) > 1 else "sub").lower()
CFG = {
 "sub": dict(src="data/subsector_flow12.json", n_expect=111,
             out="reports/SubSector_flow_watchlist_R12.00_claudeopus5high_09.18_0937.xlsx"),
 "ai":  dict(src="data/ai_flow13.json", n_expect=41,
             out="reports/AI_Sector_watchlist_R13.00_claudeopus5high_09.18_0937.xlsx"),
}[MODE]
TV = "https://www.tradingview.com/chart/Q1c5VWwD/?symbol={}"

F = json.load(open(CFG["src"])); M, ROWS = F["meta"], F["rows"]
DAYS = M["days"]
LIVE = sorted([r for r in ROWS if r.get("days")], key=lambda r: r["rank"])
V  = load_workbook(CFG["out"], data_only=True)
WB = load_workbook(CFG["out"])
P = []
def bad(m): P.append(m); print("  PROBLEM:", m)
def close(a, b, tol=1e-9):
    return a is not None and b is not None and abs(a - b) <= tol * max(1, abs(b))

print(f"[{MODE}] sheets:", V.sheetnames)
if len(ROWS) != CFG["n_expect"]: bad(f"source has {len(ROWS)} rows, expected {CFG['n_expect']}")

# ---- 總表: every live row, in rank order, with every figure matching ----
ws, wsF = V["總表 All"], WB["總表 All"]
NTK = max(len(x["basket"]) for x in LIVE)
nlink = 0
r = 9
for x in LIVE:
    d = x["days"].get(DAYS[-1])
    if ws.cell(r, 1).value != x["rank"]: bad(f"總表 row {r} rank {ws.cell(r,1).value} != {x['rank']}")
    if ws.cell(r, 2).value != x["zh"]:   bad(f"總表 row {r} name")
    if not close(ws.cell(r, 5).value, x["score5"]):    bad(f"{x['zh']} score5")
    if d and not close(ws.cell(r, 6).value, d["score"]): bad(f"{x['zh']} last-day score")
    if d and not close(ws.cell(r, 7).value, d["ret"]):  bad(f"{x['zh']} last-day ret")
    if not close(ws.cell(r, 8).value, x["ret5"]):      bad(f"{x['zh']} ret5")
    if not close(ws.cell(r, 9).value, x["breadth5"]):  bad(f"{x['zh']} breadth5")
    if not close(ws.cell(r, 10).value, x["slope"]):    bad(f"{x['zh']} slope")
    if d and not close(ws.cell(r, 11).value, d["mfd"] / 1e6, 1e-6): bad(f"{x['zh']} last-day mfd")
    if not close(ws.cell(r, 12).value, x["mfd5"] / 1e6, 1e-6):      bad(f"{x['zh']} mfd5")
    if ws.cell(r, 13).value != x["n_basket"]: bad(f"{x['zh']} n_basket")
    for j, t in enumerate(x["basket"]):
        c = wsF.cell(r, 14 + j)
        if c.value != t: bad(f"{x['zh']} basket slot {j+1} = {c.value!r} != {t!r}")
        elif c.hyperlink is None or c.hyperlink.target != TV.format(t.lower()):
            bad(f"{x['zh']} {t} link")
        else: nlink += 1
    for j in range(len(x["basket"]), NTK):
        if ws.cell(r, 14 + j).value is not None: bad(f"{x['zh']} stray basket slot {j+1}")
    r += 1
if ws.cell(r, 1).value is not None: bad("總表 has an extra row")
print(f"總表 rows verified: {len(LIVE)}, basket links: {nlink}")

# ---- 逐日分數: the 5-day grid ----
ws = V["逐日分數 Daily"]
r = 9
for x in LIVE:
    if ws.cell(r, 2).value != x["zh"]: bad(f"逐日 row {r} name")
    for i, dd in enumerate(DAYS):
        v = x["days"].get(dd)
        got = ws.cell(r, 3 + i).value
        if v and not close(got, v["score"]): bad(f"{x['zh']} {dd} daily score {got} != {v['score']}")
        if not v and got is not None: bad(f"{x['zh']} {dd} should be blank")
    if not close(ws.cell(r, 3 + len(DAYS)).value, x["score5"]): bad(f"{x['zh']} composite in 逐日")
    r += 1
print(f"逐日分數 rows verified: {len(LIVE)} x {len(DAYS)} sessions")

# ---- 摘要: top/bottom 12 must be the real top/bottom 12 ----
ws = V["摘要 Summary"]
def block_after(label):
    for rr in range(1, 80):
        if isinstance(ws.cell(rr, 1).value, str) and label in ws.cell(rr, 1).value: return rr + 2
    return None
for label, want in [("資金流入最強", LIVE[:12]), ("資金流出最強", LIVE[-12:])]:
    st = block_after(label)
    if st is None: bad(f"摘要 block {label} missing"); continue
    for i, x in enumerate(want):
        if ws.cell(st + i, 1).value != x["rank"] or ws.cell(st + i, 2).value != x["zh"]:
            bad(f"摘要 {label} row {i+1}: {ws.cell(st+i,2).value} != {x['zh']}")
print("摘要 top/bottom 12 verified")

# ---- AI-only sheets ----
if MODE == "ai":
    ws, wsF = V["成分股資金流 Constituents"], WB["成分股資金流 Constituents"]
    r, n = 9, 0
    for x in LIVE:
        for t in x["ticks"]:
            if wsF.cell(r, 3).value != t["sym"]: bad(f"constituent row {r} sym")
            elif wsF.cell(r, 3).hyperlink is None: bad(f"constituent {t['sym']} no link")
            if not close(ws.cell(r, 5).value, t["tf5"]):  bad(f"{t['sym']} tf5")
            if not close(ws.cell(r, 6).value, t["ret5"]): bad(f"{t['sym']} ret5")
            if not close(ws.cell(r, 7).value, t["mfd5"] / 1e6, 1e-6): bad(f"{t['sym']} mfd5")
            n += 1; r += 1
    exp = sum(len(x["ticks"]) for x in LIVE)
    if n != exp: bad(f"constituents {n} != {exp}")
    print(f"成分股 rows verified: {n}")
    dead = [x for x in ROWS if not x.get("days")]
    ws = V["未能計分 Unscorable"]
    for i, x in enumerate(dead):
        if ws.cell(9 + i, 2).value != x["zh"]: bad(f"unscorable row {i+1}")
    if ws.cell(9 + len(dead), 2).value is not None: bad("unscorable extra row")
    print(f"未能計分 rows verified: {len(dead)}")

# ---- 資料品質: the engine verdicts are reproduced verbatim ----
ws = V["資料品質 Data Quality"]
txt = " ".join(str(ws.cell(rr, cc).value) for rr in range(1, 40) for cc in (1, 2, 3)
               if ws.cell(rr, cc).value is not None)
for label, val in [("mirror_priceless", M.get("mirror_priceless")),
                   ("split_dropped", M.get("split_dropped")),
                   ("provisional_vol_days", M.get("provisional_vol_days"))]:
    if val and str(val) not in txt: bad(f"資料品質 does not state {label} = {val}")
if f"{M['mkt_n'][DAYS[-1]]:,}" not in txt: bad("資料品質 does not state the panel size")
print("資料品質 verdicts present")

import pandas as pd
for sh in V.sheetnames:
    if pd.read_excel(CFG["out"], sheet_name=sh, header=None).empty: bad(f"pandas read {sh} empty")
print("pandas read all sheets OK")
print(f"\n[{MODE}] PROBLEMS:", len(P))
